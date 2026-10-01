"""Выгрузка фактического выполнения программы в Excel (кнопка на дашборде админа)."""
from io import BytesIO
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import program as P

HEAD_FILL = PatternFill('solid', fgColor='0D1F3C')
HEAD_FONT = Font(bold=True, color='FFFFFF')
SUB_FILL = PatternFill('solid', fgColor='DCE3EE')
TOTAL_FONT = Font(bold=True)
THIN = Side(style='thin', color='C9D1DE')
MONEY = '#,##0'
BOTTLES = '#,##0.0'
PCT = '0%'
BRAND = {'HY': 'Hennessy', 'MC': 'Moët & Chandon'}


def _sheet(ws, headers, rows, formats, widths=None, total_row=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 42
    for r in rows:
        ws.append(r)
    if total_row:
        ws.append(total_row)
        for c in ws[ws.max_row]:
            c.font, c.fill = TOTAL_FONT, SUB_FILL
    for i, fmt in enumerate(formats, start=1):
        if not fmt:
            continue
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            row[0].number_format = fmt
    for i in range(1, len(headers) + 1):
        w = (widths or {}).get(i)
        if w is None:
            longest = max(len(str(ws.cell(row=r, column=i).value or '')) for r in range(2, min(ws.max_row, 200) + 1)) if ws.max_row > 1 else 8
            w = min(max(10, longest + 2), 42)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2' if len(headers) > 6 else 'A2'
    last = ws.max_row - (1 if total_row else 0)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(last, 1)}'


def build(rows, entries, users_by_id, closed):
    """rows — результаты outlet_summary по активным точкам; entries — все Entry."""
    wb = Workbook()

    # 1. Сводка по точкам и волнам
    ws = wb.active
    ws.title = 'Точки_волны'
    head = ['ID', 'Точка', 'Бренд', 'Город', 'Канал', 'Команда', 'Раздел', 'Ответственный',
            'План год, бут.', 'Факт год, бут.', 'Вып. год']
    fm = [None] * 8 + [BOTTLES, BOTTLES, PCT]
    for w in (1, 2, 3):
        head += [f'В{w} план', f'В{w} факт', f'В{w} вып.', f'В{w} уровень',
                 f'В{w} оборот, ₸', f'В{w} приз точке, ₸', f'В{w} бонус команде, ₸', f'В{w} итого, ₸']
        fm += [BOTTLES, BOTTLES, PCT, None, MONEY, MONEY, MONEY, MONEY]
    head += ['Заработано всего, ₸', 'Макс. при 110%, ₸', 'Ещё можно получить, ₸']
    fm += [MONEY, MONEY, MONEY]
    data = []
    tot = defaultdict(float)
    for r in sorted(rows, key=lambda r: (P.SECTION_ORDER.index(r['section']) if r['section'] in P.SECTION_ORDER else 99, -r['plan']['target'])):
        o, p = r['o'], r['plan']
        line = [o.id, p['name'], BRAND.get(p['brand'], p['brand']), p.get('city'), p.get('channel'),
                p.get('team'), r['section'], o.manager.name if o.manager else '',
                p['target'], r['bottles'], r['ann_ach']]
        for w in (1, 2, 3):
            c = r['waves'][w]
            line += [c['target'], c['bottles'], c['ach'], c['tier'], c['turnover'],
                     c['prize'], c['team'], c['total']]
            for k in ('prize', 'team', 'total', 'turnover'):
                tot[(w, k)] += c[k]
        line += [r['earned'], r['pot110'], r['left']]
        tot['earned'] += r['earned']; tot['pot110'] += r['pot110']; tot['left'] += r['left']
        data.append(line)
    total = ['', 'ИТОГО'] + [''] * 9
    for w in (1, 2, 3):
        total += ['', '', '', '', tot[(w, 'turnover')], tot[(w, 'prize')], tot[(w, 'team')], tot[(w, 'total')]]
    total += [tot['earned'], tot['pot110'], tot['left']]
    _sheet(ws, head, data, fm, widths={1: 6, 2: 30}, total_row=total)

    # 2. По месяцам
    ws = wb.create_sheet('Точки_месяцы')
    head = ['ID', 'Точка', 'Бренд', 'Команда', 'Ответственный']
    fm = [None] * 5
    for m in P.MONTHS:
        n = P.MONTH_RU[m]
        head += [f'{n} план', f'{n} факт', f'{n} вып.', f'{n} оборот, ₸']
        fm += [BOTTLES, BOTTLES, PCT, MONEY]
    data = []
    for r in rows:
        o, p = r['o'], r['plan']
        line = [o.id, p['name'], BRAND.get(p['brand'], p['brand']), p.get('team'), o.manager.name if o.manager else '']
        for m in P.MONTHS:
            c = r['months'][m]
            line += [c['target'], c['bottles'], c['ach'], c['turnover']]
        data.append(line)
    _sheet(ws, head, data, fm, widths={1: 6, 2: 30})

    # 3. По менеджерам (бонус команде идёт ответственному за точку)
    ws = wb.create_sheet('Менеджеры')
    agg = {}
    for r in rows:
        name = r['o'].manager.name if r['o'].manager else '— без ответственного —'
        a = agg.setdefault(name, dict(n=0, w={1: 0, 2: 0, 3: 0}, prize=0))
        a['n'] += 1
        for w in (1, 2, 3):
            a['w'][w] += r['waves'][w]['team']
        a['prize'] += r['earned'] - sum(r['waves'][w]['team'] for w in (1, 2, 3))
    head = ['Менеджер', 'Точек', 'Бонус В1, ₸', 'Бонус В2, ₸', 'Бонус В3, ₸', 'Бонус всего, ₸', 'Призы его точкам, ₸']
    data = [[k, v['n'], v['w'][1], v['w'][2], v['w'][3], sum(v['w'].values()), v['prize']]
            for k, v in sorted(agg.items(), key=lambda kv: -sum(kv[1]['w'].values()))]
    total = ['ИТОГО', sum(d[1] for d in data)] + [sum(d[i] for d in data) for i in range(2, 7)]
    _sheet(ws, head, data, [None, None, MONEY, MONEY, MONEY, MONEY, MONEY], widths={1: 30}, total_row=total)

    # 4. Все отгрузки (сырые данные)
    ws = wb.create_sheet('Отгрузки')
    by_id = {r['o'].id: r for r in rows}
    head = ['Дата', 'Месяц', 'Волна', 'ID', 'Точка', 'Бренд', 'SKU', 'Штук', 'Месяц закрыт', 'Внёс', 'Когда внесено']
    data = []
    for e in sorted(entries, key=lambda e: (e.date, e.outlet_id, e.sku)):
        r = by_id.get(e.outlet_id)
        if not r or not e.units:
            continue
        m = e.date[:7]
        u = users_by_id.get(e.updated_by)
        data.append([e.date, P.MONTH_RU.get(m, m), P.MONTH_WAVE.get(m, ''), e.outlet_id, r['plan']['name'],
                     BRAND.get(r['plan']['brand'], r['plan']['brand']), e.sku, e.units,
                     'да' if m in closed else 'нет', u.name if u else '',
                     e.updated_at.strftime('%Y-%m-%d %H:%M') if e.updated_at else ''])
    _sheet(ws, head, data, [None] * 7 + ['0.##'] + [None] * 3, widths={5: 30, 7: 34})

    # 5. О файле
    ws = wb.create_sheet('Пояснения')
    notes = [
        ('Выгружено', datetime.now().strftime('%d.%m.%Y %H:%M')),
        ('Закрытые месяцы', ', '.join(P.MONTH_RU[m] for m in sorted(closed)) or 'нет'),
        ('Факт', 'в бутылках-эквивалентах (Hennessy 0,7 л, Moët 0,75 л), только внесённые отгрузки'),
        ('Уровни', '<90% — без выплаты; 90–99%, 100–109%, 110%+ — ставки программы; выплата ограничена 110%'),
        ('Приз точке', 'выплачивается торговой точке через агентство'),
        ('Бонус команде', '1,5 % от зачтённого объёма по РЦ — бонус ответственного менеджера (через ФОТ)'),
        ('Открытые месяцы', 'цифры по ним могут ещё меняться, окончательные — после закрытия месяца'),
    ]
    for k, v in notes:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = TOTAL_FONT
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 100

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
